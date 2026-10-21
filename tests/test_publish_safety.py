"""
Tests for Phase 3: Publish Safety (automation/schedule_runner.py)
- Sequence order: claim -> compose -> idempotency check -> upload -> state write -> excel write
- Upload success + Excel failure -> exit code 4 and state Done
- Pipeline failure twice -> marks Failed and next row selected on 3rd call
- Rerun with video_id present -> upload mock called zero times
- Test mode never calls real upload client
- Dry run does not mark Done in Excel or state
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from unittest.mock import MagicMock, patch
from zoneinfo import ZoneInfo
import openpyxl
import pytest

import config
from automation.schedule_runner import IST, publish_next_topic
from storage.schedule_manager import ExcelLockedError
from storage import state_store


@pytest.fixture
def mock_schedule_environment(tmp_path):
    """Sets up a clean temporary Excel schedule and state file with 2 pending topics."""
    excel_path = tmp_path / "safety_schedule.xlsx"
    state_path = tmp_path / "safety_state.json"
    lock_path = tmp_path / ".test.lock"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Publishing Schedule"
    headers = [
        "Day #", "Week", "Date", "Day of Week", "Publish Time",
        "Topic", "Category", "Key Learning Objective", "Script", "Posted?"
    ]
    ws.append(headers)
    ws.append([1, 1, "20-Oct-2026", "Tuesday", "9:00 AM", "Topic One", "Macro", "Obj 1", "", "Pending"])
    ws.append([1, 1, "20-Oct-2026", "Tuesday", "3:00 PM", "Topic Two", "Macro", "Obj 2", "", "Pending"])
    wb.save(excel_path)

    state_store.load(state_file=state_path, excel_file=excel_path)

    return {
        "excel_path": excel_path,
        "state_path": state_path,
        "lock_path": lock_path,
    }


def test_publish_safety_sequence_order(mock_schedule_environment, monkeypatch, tmp_path):
    """Verifies strict execution order: claim -> compose -> idempotency check -> upload -> state write -> excel write."""
    env = mock_schedule_environment
    now_due = datetime(2026, 10, 20, 9, 30, tzinfo=IST)
    events = []

    # Mock video file
    dummy_video = tmp_path / "dummy_final.mp4"
    dummy_video.write_bytes(b"dummy video content")

    orig_claim = state_store.claim
    def spy_claim(row_key, job_id, **kwargs):
        events.append("claim")
        return orig_claim(row_key, job_id, **kwargs)

    def spy_compose(*args, **kwargs):
        events.append("compose")
        return {
            "status": "composed",
            "video_path": str(dummy_video),
            "topic": "Topic One",
            "category": "Macro",
            "script": {"full_voiceover": "voiceover text", "video_title": "Title 1"},
            "job_id": "job_seq",
        }

    from automation.schedule_runner import check_idempotency
    orig_check = check_idempotency
    def spy_idempotency_check(*args, **kwargs):
        events.append("idempotency_check")
        return orig_check(*args, **kwargs)

    def spy_upload(*args, **kwargs):
        events.append("upload")
        return {
            "video_id": "YT_SEQ_123",
            "url": "https://youtube.com/shorts/YT_SEQ_123",
            "status": "published",
        }

    orig_complete = state_store.complete
    def spy_state_write(row_key, video_id, url, **kwargs):
        events.append("state_write")
        return orig_complete(row_key, video_id, url, **kwargs)

    def spy_excel_write(row_key, expected_topic, **kwargs):
        events.append("excel_write")
        return True

    monkeypatch.setattr("storage.state_store.claim", spy_claim)
    monkeypatch.setattr("automation.schedule_runner.run_pro_pipeline", spy_compose)
    monkeypatch.setattr("automation.schedule_runner.verify_composed_video", lambda p: {"duration": 30.0})
    monkeypatch.setattr("automation.schedule_runner.check_idempotency", spy_idempotency_check)
    monkeypatch.setattr("automation.schedule_runner.publish_composed", spy_upload)
    monkeypatch.setattr("storage.state_store.complete", spy_state_write)
    monkeypatch.setattr("automation.schedule_runner.mark_row_done", spy_excel_write)

    code = publish_next_topic(
        dry_run=False,
        now=now_due,
        excel_path=env["excel_path"],
        state_file=env["state_path"],
        lock_file=env["lock_path"],
    )

    assert code == 0
    # Filter to only the core tracked lifecycle events
    lifecycle_events = [e for e in events if e in ("claim", "compose", "idempotency_check", "upload", "state_write", "excel_write")]
    assert lifecycle_events == ["claim", "compose", "idempotency_check", "upload", "state_write", "excel_write"]


def test_upload_success_excel_failure_returns_exit_4_and_state_done(mock_schedule_environment, monkeypatch, tmp_path):
    """If YouTube upload succeeds but Excel write fails, state is Done and code 4 is returned."""
    env = mock_schedule_environment
    now_due = datetime(2026, 10, 20, 9, 30, tzinfo=IST)

    dummy_video = tmp_path / "dummy.mp4"
    dummy_video.write_bytes(b"content")

    monkeypatch.setattr("automation.schedule_runner.run_pro_pipeline", lambda **k: {
        "status": "composed", "video_path": str(dummy_video), "topic": "Topic One", "script": {}
    })
    monkeypatch.setattr("automation.schedule_runner.verify_composed_video", lambda p: True)
    monkeypatch.setattr("automation.schedule_runner.publish_composed", lambda *a, **k: {
        "video_id": "YT_SUCCESS_EXCEL_FAIL", "url": "https://youtube.com/shorts/YT_SUCCESS_EXCEL_FAIL"
    })
    monkeypatch.setattr("automation.schedule_runner.mark_row_done", MagicMock(side_effect=ExcelLockedError("Excel is locked")))

    code = publish_next_topic(
        dry_run=False,
        now=now_due,
        excel_path=env["excel_path"],
        state_file=env["state_path"],
        lock_file=env["lock_path"],
    )

    assert code == 4
    # State MUST be Done with the video ID
    row = state_store.get("D01-09:00", state_file=env["state_path"])
    assert row["status"] == "Done"
    assert row["youtube_video_id"] == "YT_SUCCESS_EXCEL_FAIL"


def test_pipeline_failure_twice_marks_failed_and_selects_next_row(mock_schedule_environment, monkeypatch):
    """Pipeline failing twice marks the row Failed; the next due row is selected on the 3rd invocation."""
    env = mock_schedule_environment
    now_due = datetime(2026, 10, 20, 16, 0, tzinfo=IST)  # Both D01-09:00 and D01-15:00 are due

    # Fail during composition
    monkeypatch.setattr("automation.schedule_runner.run_pro_pipeline", MagicMock(side_effect=RuntimeError("FFmpeg crash")))

    # 1st attempt on Topic One -> fails, attempts=1, status=Pending, exit code 1
    code1 = publish_next_topic(
        dry_run=False, now=now_due, excel_path=env["excel_path"], state_file=env["state_path"], lock_file=env["lock_path"]
    )
    assert code1 == 1
    r1 = state_store.get("D01-09:00", state_file=env["state_path"])
    assert r1["attempts"] == 1
    assert r1["status"] == "Pending"

    # 2nd attempt on Topic One -> fails, attempts=2, status=Failed, exit code 1
    code2 = publish_next_topic(
        dry_run=False, now=now_due, excel_path=env["excel_path"], state_file=env["state_path"], lock_file=env["lock_path"]
    )
    assert code2 == 1
    r1 = state_store.get("D01-09:00", state_file=env["state_path"])
    assert r1["attempts"] == 2
    assert r1["status"] == "Failed"

    # 3rd invocation: Topic One is now Failed. Next due row (Topic Two D01-15:00) must be selected!
    selected_topics = []
    def spy_pipeline(topic, **kwargs):
        selected_topics.append(topic)
        raise RuntimeError("Stop after selection")

    monkeypatch.setattr("automation.schedule_runner.run_pro_pipeline", spy_pipeline)

    publish_next_topic(
        dry_run=False, now=now_due, excel_path=env["excel_path"], state_file=env["state_path"], lock_file=env["lock_path"]
    )
    assert selected_topics == ["Topic Two"]


def test_rerun_with_video_id_present_skips_upload(mock_schedule_environment, monkeypatch, tmp_path):
    """If row already has youtube_video_id when idempotency check runs, upload mock is called zero times."""
    env = mock_schedule_environment
    now_due = datetime(2026, 10, 20, 9, 30, tzinfo=IST)

    dummy_video = tmp_path / "dummy.mp4"
    dummy_video.write_bytes(b"content")

    monkeypatch.setattr("automation.schedule_runner.run_pro_pipeline", lambda **k: {
        "status": "composed", "video_path": str(dummy_video), "topic": "Topic One", "script": {}
    })
    monkeypatch.setattr("automation.schedule_runner.verify_composed_video", lambda p: True)

    upload_mock = MagicMock()
    monkeypatch.setattr("automation.schedule_runner.publish_composed", upload_mock)
    monkeypatch.setattr("automation.schedule_runner.mark_row_done", lambda *a, **k: True)

    # Pre-populate video_id in state
    state_store.upsert("D01-09:00", state_file=env["state_path"], youtube_video_id="ALREADY_UPLOADED_123")

    code = publish_next_topic(
        dry_run=False, now=now_due, excel_path=env["excel_path"], state_file=env["state_path"], lock_file=env["lock_path"]
    )

    assert code == 0
    assert upload_mock.call_count == 0


def test_test_mode_never_calls_upload_client(mock_schedule_environment, monkeypatch, tmp_path):
    """In test mode, publish_next_topic never invokes the real YouTube API upload."""
    env = mock_schedule_environment
    now_due = datetime(2026, 10, 20, 9, 30, tzinfo=IST)
    dummy_video = tmp_path / "dummy.mp4"
    dummy_video.write_bytes(b"content")

    monkeypatch.setattr(config, "PUBLISHING_MODE", "test")
    monkeypatch.setattr("automation.schedule_runner.run_pro_pipeline", lambda **k: {
        "status": "composed", "video_path": str(dummy_video), "topic": "Topic One", "script": {}, "job_id": "test_job"
    })
    monkeypatch.setattr("automation.schedule_runner.verify_composed_video", lambda p: True)
    monkeypatch.setattr("automation.schedule_runner.mark_row_done", lambda *a, **k: True)

    # If YouTubePublisher().upload is called, fail the test
    real_upload_mock = MagicMock(side_effect=AssertionError("Real YouTube upload client called in test mode!"))
    monkeypatch.setattr("publishing.youtube_publisher.YouTubePublisher.upload", real_upload_mock)

    code = publish_next_topic(
        dry_run=False, now=now_due, excel_path=env["excel_path"], state_file=env["state_path"], lock_file=env["lock_path"]
    )

    assert code == 0
    row = state_store.get("D01-09:00", state_file=env["state_path"])
    assert row["status"] == "Done"
    assert row["youtube_video_id"].startswith("TEST-")


def test_dry_run_does_not_mark_done(mock_schedule_environment, monkeypatch, tmp_path):
    """--dry-run must not mark Done in Excel or state."""
    env = mock_schedule_environment
    now_due = datetime(2026, 10, 20, 9, 30, tzinfo=IST)
    dummy_video = tmp_path / "dummy.mp4"
    dummy_video.write_bytes(b"content")

    monkeypatch.setattr("automation.schedule_runner.run_pro_pipeline", lambda **k: {
        "status": "composed", "video_path": str(dummy_video), "topic": "Topic One", "script": {}, "job_id": "dry_job"
    })

    code = publish_next_topic(
        dry_run=True, now=now_due, excel_path=env["excel_path"], state_file=env["state_path"], lock_file=env["lock_path"]
    )

    assert code == 0

    # Verify state is still Pending
    row = state_store.get("D01-09:00", state_file=env["state_path"])
    assert row["status"] == "Pending"
    assert row["youtube_video_id"] is None

    # Verify Excel is still Pending
    wb = openpyxl.load_workbook(env["excel_path"])
    assert wb.active.cell(row=2, column=10).value == "Pending"
