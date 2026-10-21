"""
End-to-end cloud publishing flow simulation tests with fake clock and mocked uploader.
Verifies the exact GitHub Actions execution sequence across a 3-slot daily schedule.
"""

import shutil
from datetime import datetime
from pathlib import Path
from unittest.mock import MagicMock, patch
from zoneinfo import ZoneInfo
import openpyxl
import pytest

import config
from storage import schedule_manager, state_store
from automation import schedule_runner

IST = ZoneInfo("Asia/Kolkata")


def test_cloud_flow_full_day_simulation(tmp_path):
    """
    Simulates three scheduled runs at 09:05, 15:20, and 20:40 IST:
    - Exactly 3 uploads produced
    - 3 distinct topics published
    - State and Excel both updated with video IDs and Done status
    - Exit code 0 on all 3 runs
    - 4th run at 20:45 exits 0 with 'nothing due'.
    """
    sim_excel = tmp_path / "cloud_schedule.xlsx"
    sim_state = tmp_path / "cloud_state.json"
    sim_lock = tmp_path / "cloud.lock"

    # Create a 3-row test schedule for Day 5 (24-Oct-2026)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Publishing Schedule"
    ws.append([
        "Day #", "Week", "Date", "Day of Week", "Publish Time",
        "Topic", "Category", "Key Learning Objective", "Script", "Posted?"
    ])
    ws.append([5, 1, "24-Oct-2026", "Saturday", "9:00 AM", "Topic Cloud 1", "Valuation", "Obj 1", "Script 1", "Pending"])
    ws.append([5, 1, "24-Oct-2026", "Saturday", "3:00 PM", "Topic Cloud 2", "Valuation", "Obj 2", "Script 2", "Pending"])
    ws.append([5, 1, "24-Oct-2026", "Saturday", "8:00 PM", "Topic Cloud 3", "Valuation", "Obj 3", "Script 3", "Pending"])
    wb.save(sim_excel)

    # Initialize state from Excel
    state_store.reconcile(excel_path=sim_excel, state_file=sim_state)

    uploaded_records: list[dict] = []

    def mock_publish(result, dry_run=False, privacy_status=None, category_id=None):
        topic = result.get("topic", "unknown")
        vid_id = f"YT-{len(uploaded_records) + 1:03d}"
        url = f"https://youtube.com/shorts/{vid_id}"
        rec = {"topic": topic, "video_id": vid_id, "url": url}
        uploaded_records.append(rec)
        return {
            "status": "published",
            "video_id": vid_id,
            "youtube_video_id": vid_id,
            "published_url": url,
            "url": url,
        }

    def mock_compose(topic, category, objective, publish_now=False, dry_run=False):
        return {
            "status": "composed",
            "topic": topic,
            "category": category,
            "video_path": str(tmp_path / "fake_video.mp4"),
            "script": {"full_voiceover": f"Full voiceover script for {topic}"},
        }

    with patch("automation.schedule_runner.publish_composed", side_effect=mock_publish), \
         patch("automation.schedule_runner.run_pro_pipeline", side_effect=mock_compose), \
         patch("automation.schedule_runner.verify_composed_video", return_value={"valid": True, "duration": 45.0}), \
         patch("automation.schedule_runner.config.validate", return_value=[]):

        # Run 1: 09:05 IST
        clock_1 = datetime(2026, 10, 24, 9, 5, tzinfo=IST)
        code_1 = schedule_runner.publish_next_topic(
            now=clock_1,
            excel_path=sim_excel,
            state_file=sim_state,
            lock_file=sim_lock,
        )
        assert code_1 == 0, f"Run 1 failed with code {code_1}"
        assert len(uploaded_records) == 1
        assert uploaded_records[0]["topic"] == "Topic Cloud 1"

        # Verify state and Excel for slot 1
        st_1 = state_store.load(state_file=sim_state)
        assert st_1["rows"]["D05-09:00"]["status"] == "Done"
        assert st_1["rows"]["D05-09:00"]["youtube_video_id"] == "YT-001"

        wb_1 = openpyxl.load_workbook(str(sim_excel), data_only=True)
        assert wb_1.active.cell(row=2, column=10).value == "Done"

        # Run 2: 15:20 IST
        clock_2 = datetime(2026, 10, 24, 15, 20, tzinfo=IST)
        code_2 = schedule_runner.publish_next_topic(
            now=clock_2,
            excel_path=sim_excel,
            state_file=sim_state,
            lock_file=sim_lock,
        )
        assert code_2 == 0, f"Run 2 failed with code {code_2}"
        assert len(uploaded_records) == 2
        assert uploaded_records[1]["topic"] == "Topic Cloud 2"

        st_2 = state_store.load(state_file=sim_state)
        assert st_2["rows"]["D05-15:00"]["status"] == "Done"
        assert st_2["rows"]["D05-15:00"]["youtube_video_id"] == "YT-002"

        wb_2 = openpyxl.load_workbook(str(sim_excel), data_only=True)
        assert wb_2.active.cell(row=3, column=10).value == "Done"

        # Run 3: 20:40 IST
        clock_3 = datetime(2026, 10, 24, 20, 40, tzinfo=IST)
        code_3 = schedule_runner.publish_next_topic(
            now=clock_3,
            excel_path=sim_excel,
            state_file=sim_state,
            lock_file=sim_lock,
        )
        assert code_3 == 0, f"Run 3 failed with code {code_3}"
        assert len(uploaded_records) == 3
        assert uploaded_records[2]["topic"] == "Topic Cloud 3"

        st_3 = state_store.load(state_file=sim_state)
        assert st_3["rows"]["D05-20:00"]["status"] == "Done"
        assert st_3["rows"]["D05-20:00"]["youtube_video_id"] == "YT-003"

        wb_3 = openpyxl.load_workbook(str(sim_excel), data_only=True)
        assert wb_3.active.cell(row=4, column=10).value == "Done"

        # Assert 3 distinct topics were uploaded
        distinct_topics = {u["topic"] for u in uploaded_records}
        assert len(distinct_topics) == 3

        # Run 4: 20:45 IST (Nothing due)
        clock_4 = datetime(2026, 10, 24, 20, 45, tzinfo=IST)
        code_4 = schedule_runner.publish_next_topic(
            now=clock_4,
            excel_path=sim_excel,
            state_file=sim_state,
            lock_file=sim_lock,
        )
        assert code_4 == 0, f"Run 4 expected exit code 0, got {code_4}"
        # Still exactly 3 uploads (no new uploads)
        assert len(uploaded_records) == 3
