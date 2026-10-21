"""
Tests for Phase 2: Schedule Data Fixes (storage/schedule_manager.py)
- Strict date and time parsing with row numbers (no silent fallbacks)
- Table-driven due-row selection (before first slot, between slots, a day late, all done)
- Real file selection: Day 2 15:00 selected when now is 2026-10-21 15:05 IST
- mark_row_done exact topic verification (refuses on mismatch)
- ExcelLockedError on PermissionError retry failure
"""

from __future__ import annotations

import shutil
from datetime import date, datetime, time as dt_time
from pathlib import Path
from unittest.mock import patch
from zoneinfo import ZoneInfo
import openpyxl
import pytest

from storage.schedule_manager import (
    IST,
    ExcelLockedError,
    _parse_date,
    _parse_time,
    mark_row_done,
    select_due_row,
)
from storage import state_store


def test_bad_time_raises_with_row_number():
    """Unparseable time raises ValueError with row number and invalid value."""
    with pytest.raises(ValueError) as exc:
        _parse_time("invalid_time_str", row_num=7)
    assert "Row 7" in str(exc.value)
    assert "invalid_time_str" in str(exc.value)

    # Valid times parse correctly
    assert _parse_time("9:00 AM") == dt_time(9, 0)
    assert _parse_time("09:00") == dt_time(9, 0)
    assert _parse_time("3:00 PM") == dt_time(15, 0)
    assert _parse_time("8:00PM") == dt_time(20, 0)
    assert _parse_time(dt_time(15, 30)) == dt_time(15, 30)


def test_bad_date_raises_with_row_number():
    """Unparseable date raises ValueError with row number and invalid value."""
    with pytest.raises(ValueError) as exc:
        _parse_date("bad_date_val", row_num=12)
    assert "Row 12" in str(exc.value)
    assert "bad_date_val" in str(exc.value)

    # Valid dates parse correctly
    assert _parse_date("20-Oct-2026") == date(2026, 10, 20)
    assert _parse_date("2026-10-21") == date(2026, 10, 21)
    assert _parse_date(date(2026, 10, 22)) == date(2026, 10, 22)


def test_due_row_selection_table_driven(tmp_path):
    """
    Table-driven due row selection:
    1. now before first slot -> None
    2. between slots -> None (next slot is upcoming)
    3. a full day late -> returns overdue row with warning
    4. all done -> None
    """
    excel_path = tmp_path / "schedule.xlsx"
    state_path = tmp_path / "state.json"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Publishing Schedule"
    headers = [
        "Day #", "Week", "Date", "Day of Week", "Publish Time",
        "Topic", "Category", "Key Learning Objective", "Script", "Posted?"
    ]
    ws.append(headers)
    ws.append([1, 1, "20-Oct-2026", "Tuesday", "9:00 AM", "Topic 1", "Macro", "Obj", "", "Pending"])
    ws.append([1, 1, "20-Oct-2026", "Tuesday", "3:00 PM", "Topic 2", "Macro", "Obj", "", "Pending"])
    wb.save(excel_path)

    # Case 1: Before first slot (20-Oct-2026 08:30 IST)
    now_before = datetime(2026, 10, 20, 8, 30, tzinfo=IST)
    row = select_due_row(now_ist=now_before, excel_path=excel_path, state_file=state_path)
    assert row is None
    assert select_due_row.reason == "nothing due"
    assert "2026-10-20T09:00:00+05:30" in str(select_due_row.next_scheduled_at)

    # Case 2: Between slots (20-Oct-2026 12:00 IST), assuming Topic 1 is marked Done in Excel and state
    wb_update = openpyxl.load_workbook(excel_path)
    wb_update.active.cell(row=2, column=10).value = "Done"
    wb_update.save(excel_path)
    state_store.complete("D01-09:00", video_id="V1", url="U1", state_file=state_path)
    now_between = datetime(2026, 10, 20, 12, 0, tzinfo=IST)
    row = select_due_row(now_ist=now_between, excel_path=excel_path, state_file=state_path)
    assert row is None
    assert "2026-10-20T15:00:00+05:30" in str(select_due_row.next_scheduled_at)

    # Case 3: A full day late (21-Oct-2026 15:00 IST), Topic 2 is still pending
    now_late = datetime(2026, 10, 21, 15, 0, tzinfo=IST)
    with patch("storage.schedule_manager.logger.warning") as mock_warn:
        row = select_due_row(now_ist=now_late, excel_path=excel_path, state_file=state_path)
        assert row is not None
        assert row["topic"] == "Topic 2"
        assert row["row_key"] == "D01-15:00"
        assert mock_warn.called
        assert "overdue" in mock_warn.call_args[0][0]

    # Case 4: All done
    wb_update = openpyxl.load_workbook(excel_path)
    wb_update.active.cell(row=3, column=10).value = "Done"
    wb_update.save(excel_path)
    state_store.complete("D01-15:00", video_id="V2", url="U2", state_file=state_path)
    now_after = datetime(2026, 10, 22, 10, 0, tzinfo=IST)
    row = select_due_row(now_ist=now_after, excel_path=excel_path, state_file=state_path)
    assert row is None


def test_select_due_row_from_copy_of_real_file(tmp_path):
    """
    Day 2 15:00 ('EV/EBITDA vs P/E') is selected from a copy of the real file
    when now is 2026-10-21 15:05 IST.
    """
    baseline_excel = Path("data/backups/publishing_schedule_baseline.xlsx")
    real_excel = baseline_excel if baseline_excel.exists() else Path("data/publishing_schedule.xlsx")
    assert real_excel.exists(), "Real publishing schedule must exist"

    copy_excel = tmp_path / "real_copy.xlsx"
    shutil.copy(real_excel, copy_excel)

    state_path = tmp_path / "state_copy.json"
    now_due = datetime(2026, 10, 21, 15, 5, 0, tzinfo=IST)

    row = select_due_row(now_ist=now_due, excel_path=copy_excel, state_file=state_path)
    assert row is not None
    assert row["row_key"] == "D02-15:00"
    assert row["topic"] == "EV/EBITDA vs P/E"
    assert row["day"] == 2
    assert row["status"] == "Pending"


def test_mark_row_done_refuses_on_topic_mismatch(tmp_path):
    """mark_row_done raises ValueError if Topic cell does not equal expected_topic."""
    excel_path = tmp_path / "mismatch_schedule.xlsx"
    state_path = tmp_path / "mismatch_state.json"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Publishing Schedule"
    ws.append(["Day #", "Week", "Date", "Day of Week", "Publish Time", "Topic", "Category", "Objective", "Script", "Posted?"])
    ws.append([1, 1, "20-Oct-2026", "Tuesday", "9:00 AM", "Real Topic Name", "Macro", "Obj", "", "Pending"])
    wb.save(excel_path)

    # Initialize state
    state_store.load(state_file=state_path, excel_file=excel_path)

    # Attempt to mark done with wrong expected topic
    with pytest.raises(ValueError) as exc:
        mark_row_done(
            row_key="D01-09:00",
            expected_topic="Completely Wrong Topic",
            excel_path=excel_path,
            state_file=state_path,
        )
    assert "Topic mismatch" in str(exc.value)

    # Verify Excel was NOT updated
    wb_check = openpyxl.load_workbook(excel_path)
    assert wb_check.active.cell(row=2, column=10).value == "Pending"


def test_mark_row_done_locked_file_simulation(tmp_path):
    """Simulated PermissionError raises ExcelLockedError after retries."""
    excel_path = tmp_path / "locked_schedule.xlsx"
    state_path = tmp_path / "locked_state.json"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Publishing Schedule"
    ws.append(["Day #", "Week", "Date", "Day of Week", "Publish Time", "Topic", "Category", "Objective", "Script", "Posted?"])
    ws.append([1, 1, "20-Oct-2026", "Tuesday", "9:00 AM", "Test Topic", "Macro", "Obj", "", "Pending"])
    wb.save(excel_path)

    state_store.load(state_file=state_path, excel_file=excel_path)

    with patch("openpyxl.Workbook.save", side_effect=PermissionError("File locked by Excel")):
        with patch("time.sleep"):  # fast test
            with pytest.raises(ExcelLockedError) as exc:
                mark_row_done(
                    row_key="D01-09:00",
                    expected_topic="Test Topic",
                    excel_path=excel_path,
                    state_file=state_path,
                )
    assert "Close" in str(exc.value)
    assert "in Excel and rerun" in str(exc.value)
