"""
Finance Shorts Pro - Reset Test Rows Utility.
Identifies any rows marked with simulated test video IDs (e.g. 'TEST-...') or test flags
and resets their state in both data/schedule_state.json and data/publishing_schedule.xlsx
back to 'Pending'.
"""

from __future__ import annotations

import argparse
import json
import logging
import os
import sys
from pathlib import Path

# Add repo root to path
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import openpyxl
import config
from storage import state_store, schedule_manager

logger = logging.getLogger(__name__)


def reset_test_rows(
    state_file: Path | None = None,
    excel_path: Path | None = None,
    dry_run: bool = False,
) -> int:
    """
    Scans state for any rows with youtube_video_id starting with 'TEST-' and reverts them
    to 'Pending' in both JSON state and the Excel schedule workbook.
    """
    s_path = state_file or (config.DATA_DIR / "schedule_state.json")
    e_path = excel_path or config.SCHEDULE_EXCEL

    if not s_path.exists():
        print(f"State file {s_path} does not exist.")
        return 0

    st = state_store.load(state_file=s_path, excel_file=e_path)
    rows = st.get("rows", {})

    test_keys = []
    for rk, r in rows.items():
        vid = str(r.get("youtube_video_id") or "")
        note = str(r.get("note") or "")
        if vid.startswith("TEST-") or "test" in note.lower():
            test_keys.append(rk)

    if not test_keys:
        print("[OK] No test rows found. Schedule state is clean.")
        return 0

    print(f"Found {len(test_keys)} test row(s) to reset: {', '.join(test_keys)}")

    if dry_run:
        print("[DRY RUN] No changes applied.")
        return 0

    # Load Excel to update cells
    wb = openpyxl.load_workbook(str(e_path))
    ws = wb["Publishing Schedule"] if "Publishing Schedule" in wb.sheetnames else wb.active

    for rk in test_keys:
        r = rows[rk]
        row_idx = r.get("row_index")
        print(f"  Resetting {rk} ('{r.get('topic')}') at Excel row {row_idx} -> Pending")

        # Reset state fields
        r["status"] = "Pending"
        r["youtube_video_id"] = None
        r["youtube_url"] = None
        r["published_at"] = None
        r["attempts"] = 0
        r["job_id"] = None
        r["claimed_at"] = None
        r["last_error"] = None
        if "note" in r and "test" in r["note"].lower():
            r.pop("note", None)

        # Reset Excel cell
        if row_idx:
            cell = ws.cell(row=row_idx, column=10)
            schedule_manager.format_status_cell(cell, "Pending", row_idx)

    # Save state
    state_store.save(st, state_file=s_path)

    # Save Excel
    schedule_manager._save_workbook_with_retry(wb, e_path)
    print(f"[SUCCESS] Successfully reset {len(test_keys)} test row(s) back to Pending.")
    return len(test_keys)


def main() -> None:
    parser = argparse.ArgumentParser(description="Reset test rows from schedule state")
    parser.add_argument("--dry-run", action="store_true", help="Preview test rows without modifying")
    parser.add_argument("--state-file", default=None, help="Path to schedule_state.json")
    parser.add_argument("--excel-file", default=None, help="Path to publishing_schedule.xlsx")

    args = parser.parse_args()
    s_path = Path(args.state_file) if args.state_file else None
    e_path = Path(args.excel_file) if args.excel_file else None

    reset_test_rows(state_file=s_path, excel_path=e_path, dry_run=args.dry_run)


if __name__ == "__main__":
    main()
