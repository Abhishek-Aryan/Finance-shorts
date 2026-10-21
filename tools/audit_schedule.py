"""
Audit script for data/publishing_schedule.xlsx.
Prints:
- total row count
- count per status (Done, Pending, etc.)
- duplicate topics
- rows whose date or time fails to parse
- rows with an empty Script among Done rows
"""

from __future__ import annotations

import sys
from collections import Counter
from datetime import date, datetime, time as dt_time
from pathlib import Path
import openpyxl


def parse_date(val: object) -> date | None:
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    if val is None:
        return None
    val_str = str(val).strip()
    for fmt in ("%d-%b-%Y", "%d %b %Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(val_str, fmt).date()
        except ValueError:
            pass
    return None


def parse_time(val: object) -> dt_time | None:
    if isinstance(val, dt_time):
        return val
    if isinstance(val, datetime):
        return val.time()
    if val is None:
        return None
    val_str = str(val).strip()
    for fmt in ("%I:%M %p", "%H:%M", "%I:%M%p"):
        try:
            return datetime.strptime(val_str, fmt).time()
        except ValueError:
            pass
    return None


def audit_schedule(excel_path: Path) -> dict:
    if not excel_path.exists():
        print(f"ERROR: Excel file not found: {excel_path}")
        sys.exit(1)

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    sheet_name = "Publishing Schedule"
    if sheet_name not in wb.sheetnames:
        ws = wb.active
    else:
        ws = wb[sheet_name]

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        print("ERROR: Empty sheet")
        sys.exit(1)

    header = rows[0]
    data_rows = rows[1:]

    status_counter = Counter()
    topic_counter = Counter()
    unparseable_dates = []
    unparseable_times = []
    done_empty_script = []

    for idx, r in enumerate(data_rows, start=2):
        if not r or all(c is None for c in r):
            continue

        day_num = r[0]
        week_num = r[1]
        date_val = r[2]
        day_of_week = r[3]
        time_val = r[4]
        topic = str(r[5]).strip() if len(r) > 5 and r[5] is not None else ""
        category = r[6] if len(r) > 6 else ""
        klo = r[7] if len(r) > 7 else ""
        script = str(r[8]).strip() if len(r) > 8 and r[8] is not None else ""
        status = str(r[9]).strip() if len(r) > 9 and r[9] is not None else "Pending"

        status_counter[status] += 1
        if topic:
            topic_counter[topic] += 1

        parsed_d = parse_date(date_val)
        if parsed_d is None:
            unparseable_dates.append((idx, date_val))

        parsed_t = parse_time(time_val)
        if parsed_t is None:
            unparseable_times.append((idx, time_val))

        if status.lower() == "done" and not script:
            done_empty_script.append((idx, topic))

    duplicate_topics = {topic: cnt for topic, cnt in topic_counter.items() if cnt > 1}

    results = {
        "total_rows": len([r for r in data_rows if any(c is not None for c in r)]),
        "status_counts": dict(status_counter),
        "duplicate_topics": duplicate_topics,
        "unparseable_dates": unparseable_dates,
        "unparseable_times": unparseable_times,
        "done_empty_script": done_empty_script,
    }

    print("=== PUBLISHING SCHEDULE AUDIT REPORT ===")
    print(f"File: {excel_path}")
    print(f"Total Data Rows: {results['total_rows']}")
    print(f"Status Counts: {results['status_counts']}")
    print(f"Duplicate Topics ({len(duplicate_topics)}): {duplicate_topics if duplicate_topics else 'None'}")
    print(f"Unparseable Dates ({len(unparseable_dates)}): {unparseable_dates if unparseable_dates else 'None'}")
    print(f"Unparseable Times ({len(unparseable_times)}): {unparseable_times if unparseable_times else 'None'}")
    print(f"Done Rows with Empty Script ({len(done_empty_script)}): {done_empty_script if done_empty_script else 'None'}")
    print("=======================================")

    return results


if __name__ == "__main__":
    target = Path("data/publishing_schedule.xlsx")
    if len(sys.argv) > 1:
        target = Path(sys.argv[1])
    audit_schedule(target)
